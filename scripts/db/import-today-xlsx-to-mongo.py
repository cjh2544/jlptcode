"""
Replace MongoDB word_today / grammar_today from Excel exports.
Deletes existing collection docs, then inserts from xlsx. Does not touch MariaDB.
"""
from __future__ import annotations

import os
from pathlib import Path

from bson import ObjectId
from openpyxl import load_workbook
from pymongo import MongoClient

ROOT = Path(__file__).resolve().parents[2]

WORD_XLSX = Path(r"C:\Users\xodus\Downloads\(261027)moku.word_today.xlsx")
GRAMMAR_XLSX = Path(r"C:\Users\xodus\Downloads\(261027)moku.grammar_today.xlsx")
CHUNK = 500


def load_env():
    for name in (".env.local", ".env", ".env.production"):
        path = ROOT / name
        if not path.exists():
            continue
        for line in path.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


load_env()


def as_str(value, fallback=None):
    if value is None or value == "":
        return fallback
    return str(value).strip()


def as_int(value, fallback=None):
    if value is None or value == "":
        return fallback
    try:
        return int(value)
    except (TypeError, ValueError):
        try:
            return int(float(value))
        except (TypeError, ValueError):
            return fallback


def as_text(value):
    if value is None or value == "":
        return None
    return str(value)


def set_if(doc: dict, key: str, value):
    if value is not None and value != "":
        doc[key] = value


def build_locale(row: dict, prefix: str):
    locale = {}
    for key in ("en", "cn", "my"):
        value = as_text(row.get(f"{prefix}.{key}"))
        if value is not None:
            locale[key] = value
    return locale or None


def build_choice_list(row: dict):
    choices = []
    any_value = False
    for idx in range(4):
        value = row.get(f"question.choice[{idx}]")
        if value is None or value == "":
            choices.append("")
        else:
            choices.append(str(value))
            any_value = True
    return choices if any_value else None


def row_map(headers, values):
    return {headers[i]: values[i] if i < len(values) else None for i in range(len(headers))}


def word_doc(row: dict):
    _id = as_int(row.get("_id"))
    if _id is None:
        return None

    doc = {
        "_id": _id,
        "level": as_str(row.get("level"), "") or "",
        "year": as_str(row.get("year"), "") or "",
        "study": as_str(row.get("study"), "") or "",
        "sortNo": as_int(row.get("sortNo"), 0) or 0,
        "wordNo": as_int(row.get("wordNo"), 0) or 0,
        "word": as_str(row.get("word"), "") or "",
        "read": as_str(row.get("read"), "") or "",
        "means": as_str(row.get("means"), "") or "",
    }

    set_if(doc, "type", as_str(row.get("type")))
    set_if(doc, "questionType", as_str(row.get("questionType")))
    set_if(doc, "sentence", as_text(row.get("sentence")))
    set_if(doc, "sentence_read", as_text(row.get("sentence_read")))
    set_if(doc, "sentence_translate", as_text(row.get("sentence_translate")))
    set_if(doc, "day", as_int(row.get("day")))
    set_if(doc, "keyword", as_str(row.get("keyword")))
    set_if(doc, "speaker", as_str(row.get("speaker")))

    question_text = as_text(row.get("question.question"))
    choices = build_choice_list(row)
    answer = as_int(row.get("question.answer"))
    if question_text is not None or choices is not None or answer is not None:
        question = {}
        if question_text is not None:
            question["question"] = question_text
        if choices is not None:
            question["choice"] = choices
        if answer is not None:
            question["answer"] = answer
        doc["question"] = question

    word_locale = build_locale(row, "word_locale")
    if word_locale:
        doc["word_locale"] = word_locale
    sentence_locale = build_locale(row, "sentence_locale")
    if sentence_locale:
        doc["sentence_locale"] = sentence_locale

    return doc


def grammar_doc(row: dict):
    raw_id = as_str(row.get("_id"))
    if not raw_id or len(raw_id) != 24:
        return None
    try:
        _id = ObjectId(raw_id)
    except Exception:
        return None

    doc = {
        "_id": _id,
        "level": as_str(row.get("level"), "") or "",
        "year": as_str(row.get("year"), "") or "",
        "study": as_str(row.get("study"), "") or "",
        "sortNo": as_int(row.get("sortNo"), 0) or 0,
    }

    set_if(doc, "sentence", as_text(row.get("sentence")))
    set_if(doc, "sentence_read", as_text(row.get("sentence_read")))
    set_if(doc, "sentence_translate", as_text(row.get("sentence_translate")))
    set_if(doc, "speaker", as_str(row.get("speaker")))
    set_if(doc, "answer", as_int(row.get("answer")))

    question_text = as_text(row.get("question.question"))
    choices = build_choice_list(row)
    if question_text is not None or choices is not None:
        question = {}
        if question_text is not None:
            question["question"] = question_text
        if choices is not None:
            question["choice"] = choices
        doc["question"] = question

    sentence_locale = build_locale(row, "sentence_locale")
    if sentence_locale:
        doc["sentence_locale"] = sentence_locale

    return doc


def load_docs(path: Path, builder):
    if not path.exists():
        raise SystemExit(f"Excel not found: {path}")
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = None
    docs = []
    skipped = 0
    for i, values in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            headers = [str(c) if c is not None else "" for c in values]
            continue
        doc = builder(row_map(headers, values))
        if doc is None:
            skipped += 1
            continue
        docs.append(doc)
    wb.close()
    return docs, skipped


def replace_collection(db, name: str, docs: list):
    col = db[name]
    before = col.count_documents({})
    col.delete_many({})
    inserted = 0
    for i in range(0, len(docs), CHUNK):
        chunk = docs[i : i + CHUNK]
        result = col.insert_many(chunk, ordered=False)
        inserted += len(result.inserted_ids)
        print(f"  {name}: {min(i + CHUNK, len(docs))}/{len(docs)}")
    after = col.count_documents({})
    by_level = list(col.aggregate([{"$group": {"_id": "$level", "n": {"$sum": 1}}}, {"$sort": {"_id": 1}}]))
    return {
        "before": before,
        "excel": len(docs),
        "inserted": inserted,
        "after": after,
        "by_level": {row["_id"]: row["n"] for row in by_level},
    }


def main():
    mongo_url = (
        os.environ.get("MONGO_COPY_URL")
        or os.environ.get("MONGODB_URL")
        or os.environ.get("MONGODB_URI")
        or ""
    ).strip()
    if not mongo_url:
        raise SystemExit("MONGO_COPY_URL (or MONGODB_URL) is required")

    word_docs, word_skipped = load_docs(WORD_XLSX, word_doc)
    grammar_docs, grammar_skipped = load_docs(GRAMMAR_XLSX, grammar_doc)
    print(f"excel word_today={len(word_docs)} skipped={word_skipped}")
    print(f"excel grammar_today={len(grammar_docs)} skipped={grammar_skipped}")

    client = MongoClient(mongo_url)
    db = client.get_default_database()
    print(f"mongo db={db.name}")

    word_stats = replace_collection(db, "word_today", word_docs)
    grammar_stats = replace_collection(db, "grammar_today", grammar_docs)

    print("word_today", word_stats)
    print("grammar_today", grammar_stats)
    client.close()


if __name__ == "__main__":
    main()
