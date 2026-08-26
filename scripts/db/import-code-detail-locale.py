"""
Update code_detail value_locale_* for strategy classification / strategyType
from the multilingual Excel export.

Usage:
  python scripts/db/import-code-detail-locale.py
  # or with a custom path:
  set CODE_DETAIL_LOCALE_XLSX=C:\\path\\to\\file.xlsx
"""
from __future__ import annotations

import os
import re
from pathlib import Path
from urllib.parse import unquote, urlparse

import pymysql
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]

CLASSIFICATION_KO_TO_KEY = {
    "문자어휘": "vocabulary",
    "문법": "grammar",
    "독해": "reading",
    "청해": "listening",
}


def load_env():
    for name in (".env.local", ".env"):
        path = ROOT / name
        if not path.exists():
            continue
        for line in path.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


load_env()

DEFAULT_XLSX = Path(r"C:\Users\xodus\Downloads\디비내용(1).xlsx")
EXCEL_PATH = Path(os.environ.get("CODE_DETAIL_LOCALE_XLSX") or DEFAULT_XLSX)


def as_str(value):
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def parse_db_url(url: str):
    parsed = urlparse(url)
    return {
        "host": parsed.hostname or "127.0.0.1",
        "port": parsed.port or 3306,
        "user": unquote(parsed.username or ""),
        "password": unquote(parsed.password or ""),
        "database": parsed.path.lstrip("/"),
        "charset": "utf8mb4",
    }


def normalize_type_key(raw):
    text = as_str(raw)
    if not text:
        return None
    text = text.replace("（", "(").replace("）", ")")
    match = re.search(r"([AB])\s*-\s*(\d+)", text, re.IGNORECASE)
    if not match:
        return None
    return f"{match.group(1).upper()}-{match.group(2)}"


def merge_locale(target: dict, ja, en, cn, my):
    for key, value in (("ja", ja), ("en", en), ("cn", cn), ("my", my)):
        text = as_str(value)
        if text and not target.get(key):
            target[key] = text


def parse_excel(path: Path):
    wb = load_workbook(path, data_only=True)
    ws = wb.active

    classification: dict[str, dict] = {}
    strategy_type: dict[str, dict] = {}

    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i == 1:
            continue
        cells = list(row) + [None] * 13
        key_raw, _b, _level, ko_cat, ko, ja_cat, ja, en_cat, en, cn_cat, cn, my_cat, my = cells[:13]

        cat_ko = as_str(ko_cat)
        if cat_ko and cat_ko in CLASSIFICATION_KO_TO_KEY:
            code_key = CLASSIFICATION_KO_TO_KEY[cat_ko]
            slot = classification.setdefault(code_key, {})
            merge_locale(slot, ja_cat, en_cat, cn_cat, my_cat)

        type_key = normalize_type_key(key_raw)
        if type_key and as_str(ko):
            slot = strategy_type.setdefault(type_key, {})
            merge_locale(slot, ja, en, cn, my)

    return classification, strategy_type


def update_rows(conn, code: str, locales: dict[str, dict]):
    updated = 0
    missing = []
    with conn.cursor() as cur:
        for key, locale in locales.items():
            cur.execute(
                """
                UPDATE `code_detail`
                   SET `value_locale_ja` = %s,
                       `value_locale_en` = %s,
                       `value_locale_cn` = %s,
                       `value_locale_my` = %s
                 WHERE `code` = %s AND `key` = %s
                """,
                (
                    locale.get("ja"),
                    locale.get("en"),
                    locale.get("cn"),
                    locale.get("my"),
                    code,
                    key,
                ),
            )
            if cur.rowcount:
                updated += cur.rowcount
            else:
                missing.append(key)
    return updated, missing


def main():
    if not EXCEL_PATH.exists():
        raise SystemExit(f"Excel not found: {EXCEL_PATH}")

    db_url = os.environ.get("DATABASE_URL")
    if not db_url:
        raise SystemExit("DATABASE_URL is required")

    classification, strategy_type = parse_excel(EXCEL_PATH)
    print(f"parsed classification={len(classification)} strategyType={len(strategy_type)}")

    conn = pymysql.connect(**parse_db_url(db_url), autocommit=False)
    try:
        c_updated, c_missing = update_rows(conn, "classification", classification)
        s_updated, s_missing = update_rows(conn, "strategyType", strategy_type)
        conn.commit()
        print(f"classification updated={c_updated} missing={c_missing}")
        print(f"strategyType updated={s_updated} missing={s_missing}")
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


if __name__ == "__main__":
    main()
